import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))
from report_data import get_report_data  # noqa: E402

OUT_DIR = ROOT / "reports" / "output"
IST = timezone(timedelta(hours=5, minutes=30))

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name=FONT, size=10)
TITLE_FONT = Font(name=FONT, bold=True, size=16, color="1F4E78")
SUB_FONT = Font(name=FONT, italic=True, size=9, color="595959")
RED_FILL = PatternFill("solid", fgColor="F8CBAD")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def write_table(ws, df, row, col=1, number_formats=None):
    """DataFrame ko header ke saath likhta hai. Aakhri data row ka number lautata hai."""
    number_formats = number_formats or {}
    for j, name in enumerate(df.columns):
        c = ws.cell(row=row, column=col + j, value=name)
        c.font, c.fill, c.border = HEADER_FONT, HEADER_FILL, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, rec in enumerate(df.itertuples(index=False), start=1):
        for j, v in enumerate(rec):
            if pd.isna(v):
                v = None
            c = ws.cell(row=row + i, column=col + j, value=v)
            c.font, c.border = BODY_FONT, BORDER
            fmt = number_formats.get(df.columns[j])
            if fmt:
                c.number_format = fmt
    return row + len(df)


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = TITLE_FONT
    ws["A2"] = sub
    ws["A2"].font = SUB_FONT


def build_latest(wb, d):
    ws = wb.create_sheet("Latest")
    df = d["latest_city"].copy()
    df.columns = ["City", "Temp (°C)", "Rain (mm)", "PM2.5 Open-Meteo",
                  "US AQI (0-500)", "US AQI category", "PM2.5 OpenWeather",
                  "OpenWeather AQI (1-5)", "OpenWeather category"]
    last = write_table(ws, df, 1, 1, {
        "Temp (°C)": "0.0", "Rain (mm)": "0.0",
        "PM2.5 Open-Meteo": "0.0", "PM2.5 OpenWeather": "0.0"})
    ws.freeze_panes = "B2"
    ws.row_dimensions[1].height = 32
    set_widths(ws, [14, 11, 10, 16, 14, 28, 18, 18, 20])
    if last >= 2:
        ws.conditional_formatting.add(f"E2:E{last}", CellIsRule(
            operator="greaterThan", formula=["150"], fill=RED_FILL))
        ws.conditional_formatting.add(f"H2:H{last}", CellIsRule(
            operator="greaterThanOrEqual", formula=["4"], fill=RED_FILL))
        ch = BarChart()
        ch.type, ch.title = "col", "PM2.5 by city: two sources (ug/m3)"
        ch.y_axis.title = "PM2.5"
        ch.add_data(Reference(ws, min_col=4, min_row=1, max_row=last), titles_from_data=True)
        ch.add_data(Reference(ws, min_col=7, min_row=1, max_row=last), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
        ch.height, ch.width = 8.5, 20
        ws.add_chart(ch, f"A{last + 3}")
        ws.cell(row=last + 2, column=1,
                value="Red cell = US AQI above 150 (Unhealthy) or OpenWeather AQI 4-5 (Poor/Very Poor). "
                      "The two AQI scales are different; do not compare them directly."
                ).font = SUB_FONT
    return last


def build_trend(wb, d):
    ws = wb.create_sheet("Trend_14d")
    df = d["trend_14d"]
    if df.empty:
        return
    p = df.pivot(index="date", columns="city", values="pm2_5_mean").reset_index()
    p.columns.name = None
    p = p.rename(columns={"date": "Date"})
    last = write_table(ws, p, 1, 1, {c: "0.0" for c in p.columns[1:]})
    ws.freeze_panes = "B2"
    set_widths(ws, [12] + [11] * (len(p.columns) - 1))
    ch = LineChart()
    ch.title, ch.height, ch.width = "Daily average PM2.5, last 14 days (ug/m3)", 9, 22
    ch.add_data(Reference(ws, min_col=2, max_col=len(p.columns), min_row=1, max_row=last),
                titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
    ws.add_chart(ch, f"A{last + 3}")


def build_forecast(wb, d):
    ws = wb.create_sheet("Forecast")
    ws["A1"] = "PM2.5 forecast, next 7 days (statistical, 80% range; NOT a guarantee)"
    ws["A1"].font = Font(name=FONT, bold=True, size=12, color="1F4E78")
    fc = d["forecast_pm25"].copy()
    fc.columns = ["City", "Date", "Forecast", "Low (80%)", "High (80%)", "Method"]
    write_table(ws, fc, 3, 1, {"Forecast": "0.0", "Low (80%)": "0.0", "High (80%)": "0.0"})
    bt = d["backtest"].copy()
    bt.columns = ["City", "Metric", "MAE model", "MAE naive baseline", "Chosen method"]
    ws.cell(row=2, column=8, value="Backtest: last 3 weeks hidden, 3 times. Lower MAE = better."
            ).font = SUB_FONT
    write_table(ws, bt, 3, 8, {"MAE model": "0.00", "MAE naive baseline": "0.00"})
    ws.freeze_panes = "A4"
    set_widths(ws, [14, 12, 10, 10, 10, 18, 3, 14, 13, 11, 18, 18])


def build_emissions(wb, d):
    ws = wb.create_sheet("Emissions")
    df = d["emissions"]
    if df.empty:
        return None
    order = [c for c in ["India", "China", "United States", "World"] if c in set(df["country"])]
    p = df.pivot(index="year", columns="country", values="co2_mt")[order].reset_index()
    p.columns.name = None
    p = p.rename(columns={"year": "Year"})
    p.columns = ["Year"] + [f"{c} CO2 (Mt)" for c in order]
    last = write_table(ws, p, 1, 1, {c: "#,##0" for c in p.columns[1:]})
    n = len(order)
    for j, name in enumerate(order):                      # YoY growth (formula)
        col = 2 + n + j
        h = ws.cell(row=1, column=col, value=f"{name} YoY %")
        h.font, h.fill, h.border = HEADER_FONT, HEADER_FILL, BORDER
        h.alignment = Alignment(horizontal="center", wrap_text=True)
        src = get_column_letter(2 + j)
        for r in range(2, last + 1):
            c = ws.cell(row=r, column=col, value=f'=IFERROR({src}{r}/{src}{r-1}-1,"")')
            c.number_format, c.font, c.border = "0.0%", BODY_FONT, BORDER
    share_col = 2 + 2 * n
    if "India" in order and "World" in order:             # India ka hissa (formula)
        h = ws.cell(row=1, column=share_col, value="India share of world %")
        h.font, h.fill, h.border = HEADER_FONT, HEADER_FILL, BORDER
        h.alignment = Alignment(horizontal="center", wrap_text=True)
        a, b = get_column_letter(2 + order.index("India")), get_column_letter(2 + order.index("World"))
        for r in range(2, last + 1):
            c = ws.cell(row=r, column=share_col, value=f"=IFERROR({a}{r}/{b}{r},\"\")")
            c.number_format, c.font, c.border = "0.0%", BODY_FONT, BORDER
    ws.row_dimensions[1].height = 32
    set_widths(ws, [8] + [16] * (2 * n + 1))
    ws.cell(row=last + 2, column=1,
            value="Source: Our World in Data CO2 dataset (annual, fossil CO2 + cement, country level, "
                  "million tonnes). Not live data."
            ).font = SUB_FONT
    ch = LineChart()
    ch.title, ch.height, ch.width = "Annual CO2 emissions (Mt)", 8.5, 18
    ch.add_data(Reference(ws, min_col=2, max_col=1 + n, min_row=1, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
    ws.add_chart(ch, f"A{last + 4}")
    india_col = get_column_letter(2 + order.index("India")) if "India" in order else None
    india_yoy = get_column_letter(2 + n + order.index("India")) if "India" in order else None
    return {"last": last, "india_col": india_col, "india_yoy": india_yoy}


def build_summary(wb, d, latest_last, em, now):
    ws = wb["Summary"]
    title(ws, "Climate Impact Dashboard: Daily MIS",
          f"Generated {now:%d %b %Y, %H:%M} IST  |  Cities: Delhi, Mumbai, Kolkata, Chennai, "
          "Bengaluru, Hyderabad, Pune, Ahmedabad, Lucknow, Jaipur")
    hdr = ["KPI", "Value", "Note"]
    for j, h in enumerate(hdr, start=1):
        c = ws.cell(row=4, column=j, value=h)
        c.font, c.fill, c.border = HEADER_FONT, HEADER_FILL, BORDER
    L = latest_last
    rows = [
        ("Cities tracked", f"=COUNTA(Latest!A2:A{L})", None, "Cities in latest snapshot"),
        ("Average temperature (°C)", f"=AVERAGE(Latest!B2:B{L})", "0.0", "OpenWeather live reading"),
        ("Hottest city", f"=INDEX(Latest!A2:A{L},MATCH(MAX(Latest!B2:B{L}),Latest!B2:B{L},0))", None, "Right now"),
        ("Highest PM2.5, Open-Meteo (ug/m3)", f"=MAX(Latest!D2:D{L})", "0.0", "Modelled air quality"),
        ("City with highest PM2.5", f"=INDEX(Latest!A2:A{L},MATCH(MAX(Latest!D2:D{L}),Latest!D2:D{L},0))", None, "Open-Meteo"),
        ("Cities with US AQI above 150", f'=COUNTIF(Latest!E2:E{L},">150")', "0", "Unhealthy or worse"),
        ("Cities with OpenWeather AQI 4-5", f'=COUNTIF(Latest!H2:H{L},">=4")', "0", "Poor or Very Poor"),
    ]
    if em and em["india_col"]:
        r = em["last"]
        rows += [
            ("India CO2, latest year (Mt)", f"=Emissions!{em['india_col']}{r}", "#,##0", "Annual, country level"),
            ("India CO2 growth vs previous year", f"=Emissions!{em['india_yoy']}{r}", "0.0%", "Annual, not live"),
        ]
    for i, (k, f, fmt, note) in enumerate(rows, start=5):
        for j, v in enumerate((k, f, note), start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font, c.border = BODY_FONT, BORDER
            if j == 2 and fmt:
                c.number_format = fmt
            if j == 2:
                c.alignment = Alignment(horizontal="right")
    end = 5 + len(rows) + 1
    notes = [
        "How to read: every number in Value is a formula that reads the Latest and Emissions sheets.",
        "Air quality is modelled data (Open-Meteo CAMS, OpenWeather), not ground-sensor data.",
        "Forecasts are statistical short-term trends. They cannot predict seasonal spikes.",
    ]
    for k, t in enumerate(notes):
        ws.cell(row=end + k, column=1, value=t).font = SUB_FONT
    set_widths(ws, [38, 18, 38])


def main():
    now = datetime.now(IST)
    data = get_report_data()
    wb = Workbook()
    wb.active.title = "Summary"
    latest_last = build_latest(wb, data)
    build_trend(wb, data)
    build_forecast(wb, data)
    em = build_emissions(wb, data)
    build_summary(wb, data, max(latest_last, 2), em, now)
    wb.move_sheet("Summary", offset=-len(wb.sheetnames) + 1)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / f"MIS_{now:%Y-%m-%d}.xlsx"
    wb.save(path)
    print("Excel MIS bani:", path)
    return path


if __name__ == "__main__":
    main()